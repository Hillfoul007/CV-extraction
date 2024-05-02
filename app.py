
from flask import Flask, render_template, request, send_file, send_from_directory, redirect, url_for
from werkzeug.utils import secure_filename
import os
import re
from docx import Document
import PyPDF2
import xlwt
from xlwt import Workbook

app = Flask(__name__)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'docx', 'pdf'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

emails_list = []
phones_list = []
overall_text_list = []


def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS


def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    full_text = ""
    for para in doc.paragraphs:
        full_text += para.text + "\n"
    return full_text


def extract_text_from_pdf(pdf_file):
    reader = PyPDF2.PdfReader(pdf_file)
    full_text = ""
    for page in reader.pages:
        full_text += page.extract_text() + "\n"
    return full_text


def extract_information(cv_text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b' or 'Email -id:'
    phone_pattern = r'\b(?:\+?(\d{1,3}))?[-. (]*(\d{3})[-. )]*(\d{3})[-. ]*(\d{4})(?: *x(\d+))?\b'
    emails = re.findall(email_pattern, cv_text)
    phones = re.findall(phone_pattern, cv_text)
    return emails, phones, cv_text


def write_to_excel(emails_list, phones_list, overall_text_list):
    wb = Workbook()
    sheet = wb.add_sheet('CV Information')
    sheet.write(0, 0, 'Email ID')
    sheet.write(0, 1, 'Contact Number')
    sheet.write(0, 2, 'Overall Text')

    row = 1
    for i in range(len(emails_list)):
        emails = emails_list[i]
        phones = phones_list[i] if i < len(phones_list) else []
        overall_text = overall_text_list[i]

        # Determine the number of rows needed for this file
        num_rows = max(len(emails), len(phones), 1)  # Ensure at least one row is written

        for j in range(num_rows):
            if j < len(emails):
                sheet.write(row, 0, emails[j])
            if j < len(phones):
                sheet.write(row, 1, phones[j])
            sheet.write(row, 2, overall_text)
            row += 1

    # Save the Excel file to the uploads folder
    filename = 'CV_Information.xls'
    excel_file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
    wb.save(excel_file_path)
    return excel_file_path



@app.route('/')
def index():
    return render_template('Upload.html', excel_file=None)


@app.route('/upload', methods=['POST'])
def upload_file():
    if request.method == 'POST':
        # Reset lists for each new request
        emails_list.clear()
        phones_list.clear()
        overall_text_list.clear()
        successful_files = []
        failed_files = []

        files = request.files.getlist('file')
        for file in files:
            if file.filename == '':
                return 'No selected file'
            if file and allowed_file(file.filename):
                filename = secure_filename(file.filename)
                file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
                file.save(file_path)

                if filename.endswith('.docx'):
                    cv_text = extract_text_from_docx(file_path)
                elif filename.endswith('.pdf'):
                    cv_text = extract_text_from_pdf(file_path)
                else:
                    failed_files.append(filename)
                    continue

                emails, phones, overall_text = extract_information(cv_text)
                emails_list.append(emails)
                phones_list.append(phones)
                overall_text_list.append(overall_text)
                successful_files.append(filename)
            else:
                failed_files.append(filename)

        # Write information to Excel
        excel_file_path = write_to_excel(emails_list, phones_list, overall_text_list)
        if failed_files:
            error_message = f"Failed to process the following files: {', '.join(failed_files)}"
            return render_template('Upload.html', excel_file=excel_file_path, error_message=error_message)

        return redirect(url_for('index'))


from flask import send_from_directory

@app.route('/download_excel')
def download_excel():
    filename = 'CV_Information.xls'
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)

